#!/usr/bin/env python3
"""Build data/splits/city/aviation.json: Chicago Department of Aviation and airport fund detail.

Run:  python3 scripts/aviation_build.py            (uses cached raw/aviation/*, fetches what is missing)
      python3 scripts/aviation_build.py --refresh  (re-pull FAA awards from USASpending.gov)
Then: python3 build/city_tree.py   (every split must show applied, 0 skipped)

What it writes (research/aviation.md has the sources, checks and what was NOT attached and why):

 1. 925F-2810-909A  O'Hare AIP reserve $411,790,000  -> named FAA grants (FY2025 and earlier) with unspent money.
 2. 925F-2805-909A  Midway AIP reserve $123,869,000  -> same.
 3. 925F-2810-0140  O'Hare AIP professional services $112,740,000 -> paid_to_date: Terminal 3 construction
    manager contract 234140 (its own text says "Federal Funds" and "O'Hare"). FY2026 FAA awards go in as side info.
 4. 925F-2805-0140  Midway AIP professional services $36,199,000 -> the FY2026 Midway AIP awards (ALN 20.116).
 5. O'Hare interest "Other bonds" residual -> Series 2010B Build America Bonds interest, computed from printed
    par and coupons and checked against the printed federal subsidy.
 6. Finance General 0740 0140 ($92.2M) and 9047 (special capital projects, $2.0M): side info only (prior-year
    budget book numbers, O'Hare capital program tables from the 2026CD official statement).

Lines that already carry data/splits/city/paid_to_date.json (0740-2015-0140 and 0610-2010-0140) are NOT touched here:
a second split on the same box would be skipped (SPLITS.md: a target must be a leaf).
Nothing is scaled. Every piece fits inside its line. Never invents numbers.
"""
import csv
import json
import os
import re
import subprocess
import sys
import time
from collections import defaultdict
from datetime import datetime

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
RAW = os.path.join(ROOT, "raw", "aviation")
OUT = os.path.join(ROOT, "data", "splits", "city", "aviation.json")
AWARDS_CACHE = os.path.join(RAW, "faa_awards_usaspending.json")
os.makedirs(RAW, exist_ok=True)

USA = "https://api.usaspending.gov/api/v2"
ZIP_AIRPORT = {"60666": "ORD", "60638": "MDW"}
AIP_ALN = {"20.106", "20.116"}

# ---------------------------------------------------------------- ordinance line amounts (dataset 6694-f78c)
ORD_ORD_RESERVE = 411_790_000
ORD_MDW_RESERVE = 123_869_000
ORD_ORD_0140 = 112_740_000
ORD_MDW_0140 = 36_199_000
ORD_FG_0140 = 92_222_410
ORD_FG_9047 = 2_000_000
ORD_OHARE_INTEREST_RESIDUAL = 79_662_262   # 0902 interest line minus the 31 printed series (city_tree.py)

SRC_ORD = {"dataset": "6694-f78c", "name": "2026 Annual Appropriation Ordinance",
           "url": "https://data.cityofchicago.org/Administration-Finance/Budget-2026-Budget-Ordinance-Appropriations/6694-f78c"}
SRC_BOOK = {"doc": "Mayor's Budget Recommendations for Year 2026",
            "url": "https://www.chicago.gov/content/dam/city/depts/obm/supp_info/2026Budget/2026%20Budget%20Recommendations.pdf"}
SRC_FAA_HIST = {"doc": "FAA Airport Improvement Program grant histories",
                "url": "https://www.faa.gov/airports/aip/grant_histories"}
SRC_OS = {"doc": "City of Chicago O'Hare General Airport Senior Lien Revenue Refunding Bonds 2026C and 2026D, Official Statement dated 2026-08-21",
          "file": "raw/bonds/ohare_2026CD_OS.pdf"}


# ---------------------------------------------------------------- USASpending (federal award records)
def curl_json(url, body=None):
    cmd = ["curl", "-s", "-m", "90", "-A", "Mozilla/5.0", url]
    if body is not None:
        cmd += ["-H", "Content-Type: application/json", "-d", json.dumps(body)]
    for _ in range(3):
        r = subprocess.run(cmd, capture_output=True, text=True)
        try:
            return json.loads(r.stdout)
        except Exception:
            time.sleep(2)
    raise SystemExit("failed " + url)


def fetch_awards():
    """All FAA assistance awards to the City of Chicago since FAA FY2021, with detail (obligation, outlay, place, period)."""
    ids = {}
    for page in range(1, 20):
        body = {"filters": {"award_type_codes": ["02", "03", "04", "05"],
                            "time_period": [{"start_date": "2020-10-01", "end_date": datetime.now().strftime("%Y-%m-%d"),
                                             "date_type": "action_date"}],
                            "recipient_search_text": ["CITY OF CHICAGO"],
                            "agencies": [{"type": "awarding", "tier": "subtier", "name": "Federal Aviation Administration"}]},
                "fields": ["Award ID", "Recipient Name", "Award Amount"], "page": page, "limit": 100,
                "sort": "Award Amount", "order": "desc", "subawards": False}
        r = curl_json(USA + "/search/spending_by_award/", body)
        for x in r.get("results", []):
            ids[x["Award ID"]] = x
        if not r.get("page_metadata", {}).get("hasNext"):
            break
    # also keep every award id already known from the earlier pull in raw/grants (so nothing disappears)
    old = os.path.join(ROOT, "raw", "grants", "faa_awards_detail.json")
    if os.path.exists(old):
        for x in json.load(open(old)):
            ids.setdefault(x["award_id"], {"Award ID": x["award_id"]})
    out = []
    for aid in sorted(ids):
        gid = None
        r = curl_json(USA + "/search/spending_by_award/",
                      {"filters": {"award_type_codes": ["02", "03", "04", "05"], "award_ids": [aid]},
                       "fields": ["Award ID"], "limit": 2})
        if r.get("results"):
            gid = r["results"][0]["generated_internal_id"]
        if not gid:
            continue
        d = curl_json(USA + "/awards/" + gid + "/")
        pop = (d.get("place_of_performance") or {})
        out.append({"award_id": aid, "generated_internal_id": gid, "recipient": (d.get("recipient") or {}).get("recipient_name"),
                    "total_obligation": d.get("total_obligation"), "total_outlay": d.get("total_outlay"),
                    "description": d.get("description"), "zip5": pop.get("zip5"),
                    "start": (d.get("period_of_performance") or {}).get("start_date"),
                    "end": (d.get("period_of_performance") or {}).get("end_date"),
                    "aln": [c.get("cfda_number") for c in (d.get("cfda_info") or [])],
                    "usaspending_url": "https://www.usaspending.gov/award/" + gid})
        time.sleep(0.15)
    json.dump({"pulled": datetime.now().strftime("%Y-%m-%d %H:%M"), "awards": out}, open(AWARDS_CACHE, "w"), indent=1)
    return out


def load_awards(refresh=False):
    if refresh or not os.path.exists(AWARDS_CACHE):
        return fetch_awards()
    return json.load(open(AWARDS_CACHE))["awards"]


# ---------------------------------------------------------------- FAA grant histories (grant number, project summary)
def faa_workbook_index():
    """FAIN -> {grant_number, summary, amount, award_date, airport} for ORD and MDW rows in the cached FAA workbooks."""
    import warnings
    import openpyxl
    warnings.filterwarnings("ignore")
    idx = {}
    for fn in sorted(os.listdir(RAW)):
        if not (fn.startswith("FY") and fn.endswith(".xlsx")) or "AIP" not in fn.upper():
            continue
        wb = openpyxl.load_workbook(os.path.join(RAW, fn), read_only=True, data_only=True)
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                cells = [c for c in r if c is not None]
                if not any(isinstance(c, str) and re.fullmatch(r"3-17-00(22|25)-\d{3}-\d{4}", c) for c in cells):
                    continue
                gn = next(c for c in cells if isinstance(c, str) and re.fullmatch(r"3-17-00(22|25)-\d{3}-\d{4}", c))
                nums = [c for c in cells if isinstance(c, (int, float))]
                date = next((c for c in cells if isinstance(c, datetime)), None)
                summ = cells[-1] if isinstance(cells[-1], str) and not re.fullmatch(r"3-17-.*", cells[-1]) else ""
                fain = gn.replace("-", "")
                idx[fain] = {"grant_number": gn, "summary": summ, "amount": nums[-1] if nums else None,
                             "award_date": date.strftime("%Y-%m-%d") if date else None, "file": fn}
    return idx


def project_name(summary, desc):
    """Short plain name: FAA summary (deduplicated), else the first PURPOSE phrase of the USASpending text."""
    if summary:
        parts = []
        for p in re.split(r",\s*", summary):
            p = p.strip()
            if p and p.lower() not in [q.lower() for q in parts]:
                parts.append(p)
        return "; ".join(parts)
    m = re.match(r"PURPOSE:\s*([^.]+)\.", desc or "")
    return (m.group(1).strip().capitalize() if m else (desc or "FAA grant")[:60])


def fy_of(date_str):
    """FAA/federal fiscal year of a yyyy-mm-dd date (starts Oct 1)."""
    y, m = int(date_str[:4]), int(date_str[5:7])
    return y + 1 if m >= 10 else y


def r2(x):
    return round(float(x), 2)


WHY_AWARD = ("The FAA has promised this money to the City for this one airport project. It is paid back to the City as work "
             "is done, and no public list splits the project into smaller pieces.")


def award_pieces(awards, idx, airport, fy_pred, label_for_note):
    pieces, total, audit = [], 0.0, []
    for a in awards:
        if ZIP_AIRPORT.get(a["zip5"]) != airport:
            continue
        if not fy_pred(fy_of(a["start"])):
            continue
        if not (set(a.get("aln") or []) & AIP_ALN):
            continue   # only the Airport Improvement Program (20.106 umbrella, 20.116). 20.117 (IIJA grants) and 20.931 are other budget lines
        ob = a["total_obligation"] or 0.0
        out = a["total_outlay"]
        unspent = ob - (out or 0.0)
        if unspent <= 1 or (a["end"] or "") < "2026-01-01":
            continue
        f = idx.get(a["award_id"])
        gn = f["grant_number"] if f else None
        # cross-check the FAA workbook amount against the USASpending obligation
        if f and f["amount"] is not None and abs(f["amount"] - ob) > 1:
            audit.append(f"{gn}: FAA workbook {f['amount']:,.0f} vs USASpending {ob:,.0f}")
        pname = project_name(f["summary"] if f else "", a["description"])
        nm = f"{pname} (FAA grant {gn})" if gn else f"{pname} (FAA grant {a['award_id']})"
        note = (f"Award {a['award_id']}: {ob:,.2f} awarded, {(out or 0):,.2f} paid out so far"
                + (" (no payout reported yet, counted as zero)" if out is None else "")
                + f", period {a['start']} to {a['end']}. The box shows the part not yet paid out. "
                + ("Amount and project type match the FAA grant history. " if f else "Not yet in an FAA grant history file (the FAA has not published its FY2026 list), so USASpending is the only source. ")
                + label_for_note)
        pieces.append({"name": nm, "amount": r2(unspent), "basis": "gov_estimate", "kind": "grant_award",
                       "source": {"doc": "USASpending.gov award record", "url": a["usaspending_url"],
                                  "faa_grant_history": f["file"] if f else None},
                       "why": WHY_AWARD, "note": note,
                       "extra": {"award_id": a["award_id"], "faa_grant_number": gn, "obligation": r2(ob),
                                 "outlay": r2(out) if out is not None else None, "period_start": a["start"], "period_end": a["end"]}})
        total += r2(unspent)
    pieces.sort(key=lambda p: -p["amount"])
    return pieces, round(total, 2), audit


def award_side(awards, airport, fy_pred, kind, label):
    rows = []
    for a in awards:
        if ZIP_AIRPORT.get(a["zip5"]) != airport or not fy_pred(fy_of(a["start"])):
            continue
        rows.append({"kind": kind, "label": f"{label}: {project_name('', a['description'])} (award {a['award_id']}, ALN {', '.join(a.get('aln') or [])})",
                     "amount": r2(a["total_obligation"]), "period": f"{a['start']} to {a['end']}", "basis": "actual",
                     "source": {"doc": "USASpending.gov award record", "url": a["usaspending_url"]}})
    rows.sort(key=lambda s: -s["amount"])
    return rows


# ---------------------------------------------------------------- PFC (FAA approved locations workbook)
def pfc_side():
    import warnings
    import openpyxl
    warnings.filterwarnings("ignore")
    p = os.path.join(RAW, "pfc_locations_latest.xlsx")
    if not os.path.exists(p):
        return {}
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    out = {}
    for r in wb.worksheets[0].iter_rows(values_only=True):
        if r and r[3] in ("ORD", "MDW") and r[8]:
            out[r[3]] = {"approved": r[8], "level": r[5], "start": r[6].strftime("%Y-%m-%d"), "expires": r[7].strftime("%Y-%m-%d")}
    return out


# ---------------------------------------------------------------- Terminal 3 construction manager payments (contract 234140)
def t3_payments():
    rows = [r for r in csv.DictReader(open(os.path.join(ROOT, "raw/contracts/payments_2026ytd_dedup.csv")))
            if r["contract_number"] == "234140" and r["department_name"] == "CHICAGO DEPARTMENT OF AVIATION"]
    by_month = defaultdict(lambda: [0.0, 0])
    for r in rows:
        m, d, y = r["check_date"].split("/")
        by_month[f"{y}-{m}"][0] += float(r["amount"])
        by_month[f"{y}-{m}"][1] += 1
    return rows, by_month


MONTHS = {"01": "January", "02": "February", "03": "March", "04": "April", "05": "May", "06": "June", "07": "July",
          "08": "August", "09": "September", "10": "October", "11": "November", "12": "December"}


def main():
    refresh = "--refresh" in sys.argv
    awards = load_awards(refresh)
    idx = faa_workbook_index()
    pfc = pfc_side()
    splits = []
    audits = []

    # ---- 1 and 2: AIP reserve lines (carryover): FY2025 and earlier awards with money still unspent
    for airport, amount, auth, nm in (("ORD", ORD_ORD_RESERVE, "2810", "O'Hare"), ("MDW", ORD_MDW_RESERVE, "2805", "Midway")):
        pcs, tot, au = award_pieces(
            awards, idx, airport, lambda fy: fy <= 2025,
            "Matched to the carryover because the award was made in FAA fiscal year 2025 or earlier. The ordinance figure is a "
            "carryover estimate, so this is an upper bound of what the named awards can explain.")
        audits += au
        assert tot <= amount, (airport, tot, amount)
        side = [{"kind": "grant_carryover", "label": "Summary G: carryover from 2025 (Aug 1 2025 estimate)", "amount": amount,
                 "period": "2026 budget", "basis": "published", "source": dict(SRC_BOOK, page="607 (PDF 616)")},
                {"kind": "grant_anticipated_2026", "label": "Summary G: new grant money anticipated in 2026 (the 0140 line below)",
                 "amount": ORD_ORD_0140 if airport == "ORD" else ORD_MDW_0140, "period": "2026 budget", "basis": "published",
                 "source": dict(SRC_BOOK, page="607 (PDF 616)")}]
        if airport in pfc:
            p = pfc[airport]
            side.append({"kind": "pfc_authority", "label": f"FAA passenger facility charge authority at {nm}: ${p['level']:.2f} per passenger from {p['start']} to {p['expires']} (total approved, all projects)",
                         "amount": p["approved"], "period": "as of 2026-09-30", "basis": "published",
                         "source": {"doc": "FAA PFC approved locations, collections and expiration dates", "url": "https://www.faa.gov/airports/pfc/monthly_reports",
                                    "file": "raw/aviation/pfc_locations_latest.xlsx",
                                    "note": "The FAA publishes one total per airport. No project list with amounts was found, so PFC money is not split into boxes."}})
        if airport == "ORD":
            for lab, amt in (("Total airport capital program, escalated dollars (2026CD OS Table S-1)", 14_627_114_000),
                             ("  of which already funded by earlier bonds and about $290.1M of executed grants", 5_537_774_000),
                             ("  of which paid by the 2026B bonds", 1_004_658_000),
                             ("  of which to be paid by future bonds (no future grants assumed)", 7_484_681_000),
                             ("  of which paid from airline fees (pre-approved allowances)", 600_000_000)):
                side.append({"kind": "capital_program", "label": lab, "amount": amt, "period": "program to 2035", "basis": "published",
                             "source": dict(SRC_OS, page="E-19 and E-46 (Tables S-1, 2-2)")})
        splits.append({
            "target": {"by": "ordinance_line", "fund": "925F", "dept": "85", "authority": auth, "account": "909A"},
            "expect_amount": amount, "mode": "budget_split", "pieces": pcs,
            "residual": {"name": "Carryover not tied to a named FAA award",
                         "why": ("The budget carries this much federal airport grant money over from earlier years. The FAA and USASpending records name "
                                 "awards that still have unspent money, but they explain only part of it. The rest is carryover the budget expects "
                                 "and no public record ties it to a project.")},
            "source": SRC_ORD,
            "note": (f"Named pieces: {len(pcs)} FAA grants made in FAA fiscal year 2025 or earlier that still have unspent money "
                     f"(${tot:,.2f} in all) per USASpending.gov, cross-checked with the FAA AIP grant histories (grant numbers and project types). "
                     "Awards made after September 2025 are shown on the 'professional and technical services' line of the same grant, because they are the 2026 new money. "
                     "The pieces are an upper bound: the FAA does not publish which grant pays which budget line."),
            "side": side})

    # ---- 4: Midway 925F 0140: FY2026 awards (the 2026 anticipated grant money)
    pcs, tot, au = award_pieces(
        awards, idx, "MDW", lambda fy: fy >= 2026,
        "Matched to this line because the award was made in FAA fiscal year 2026 and this line is the 2026 anticipated grant (Summary G).")
    audits += au
    if pcs and tot <= ORD_MDW_0140:
        splits.append({
            "target": {"by": "ordinance_line", "fund": "925F", "dept": "85", "authority": "2805", "account": "0140"},
            "expect_amount": ORD_MDW_0140, "mode": "budget_split", "pieces": pcs,
            "residual": {"name": "2026 grant money not yet awarded",
                         "why": ("The budget expects this much new federal grant money in 2026. The FAA had not announced awards for all of it "
                                 "when the data was collected (October 2026).")},
            "source": SRC_ORD,
            "note": (f"{len(pcs)} Midway FAA awards made in August and September 2026 (${tot:,.2f}) against the ${ORD_MDW_0140:,.0f} the budget "
                     "anticipated as new 2026 grant money (Mayor's Budget Recommendations 2026, Summary G p.607). The match is by year, not by an official "
                     "crosswalk. No 2026 payment could be matched to this line without doubt (the Midway runway contract text does not say it is federally funded).")})

    # ---- 3: O'Hare 925F 0140: paid to date on contract 234140 + FY2026 awards as side info
    rows, by_month = t3_payments()
    paid = round(sum(r2(v[0]) for v in by_month.values()), 2)   # months rounded first so children add up to the piece in cents
    assert paid <= ORD_ORD_0140
    assert abs(paid - sum(float(r["amount"]) for r in rows)) < 0.05
    kids = []
    for ym in sorted(by_month):
        amt, n = by_month[ym]
        kids.append({"name": f"{MONTHS[ym[5:]]} {ym[:4]}", "amount": r2(amt), "basis": "paid_to_date", "kind": "month",
                     "note": f"{n} payments by check date",
                     "why": "These are checks the City wrote to this one company in this month. The payment records show each check, not what it bought."})
    awards_fy26 = award_side(awards, "ORD", lambda fy: fy >= 2026, "faa_award_fy2026", "FAA award made in FAA fiscal year 2026")
    splits.append({
        "target": {"by": "ordinance_line", "fund": "925F", "dept": "85", "authority": "2810", "account": "0140"},
        "expect_amount": ORD_ORD_0140, "mode": "paid_to_date",
        "pieces": [{"name": "Clark-W.E. O'Neil JV: construction manager for Terminal 3 (contract 234140)", "amount": paid,
                    "basis": "paid_to_date", "kind": "vendor",
                    "note": ("Contract 234140, Specification 1258275: Construction Management At-Risk Services at Terminal 3, O'Hare International Airport, "
                             f"'Federal Funds'. {len(rows)} checks from 01/12/2026 to 09/17/2026. Matched to this line because the contract text names O'Hare and federal funds, "
                             "its type is professional services, and no other O'Hare federal contract of that type was paid. FAA terminal-expansion grants for Terminal 3 are listed on the "
                             "reserve line of this grant."),
                    "extra": {"contract": "234140", "payments": len(rows)},
                    "why": "This is what one company has been paid so far this year to manage building work at Terminal 3, and the payment records do not split it by task.",
                    "children": kids}],
        "source": {"dataset": "s4vu-giwb", "name": "City of Chicago Payments (deduplicated, checks Jan 1 to 09/28/2026)",
                   "url": "https://data.cityofchicago.org/d/s4vu-giwb", "note": "joined to Contracts rsxa-ify5 on contract number; research/aviation.md"},
        "residual": {"name": "Budgeted but not spent yet",
                     "why": "This is grant budget the City has set aside for this line but has not paid out yet this year (payments run to 09/28/2026)."},
        "note": ("Only payments that match this line without doubt are shown. Other O'Hare federally funded contracts were not placed: two construction contracts "
                 "(Paschen, 21.9M together) could be charged here or to the carryover line, and the airport engineering task orders do not name an airport. "
                 "The side list shows the FAA awards made in FAA fiscal year 2026, which is what this line anticipated as new 2026 grant money."),
        "side": awards_fy26})

    # ---- 5: O'Hare interest residual: Series 2010B BABs
    t1, t2 = 315_500_000, 12_500_000
    i1, i2 = round(t1 * 0.06395, 2), round(t2 * 0.06145, 2)
    interest_2010b = round(i1 + i2, 2)
    subsidy_check = 0.35 * interest_2010b * (1 - 0.057)
    assert abs(subsidy_check - 6_912_000) < 1_000, subsidy_check   # printed: $6,912 thousand (2026CD OS Table B-4)
    splits.append({
        "target": {"by": "id", "id": "city.loans.chicago-o-hare-airport-fund.0740-2005-0902.other-not-itemised"},
        "expect_amount": ORD_OHARE_INTEREST_RESIDUAL, "mode": "budget_split",
        "pieces": [{"name": "Series 2010B (Build America Bonds, taxable)", "amount": interest_2010b, "basis": "tied", "kind": "bond_series",
                    "source": dict(SRC_OS, page="p.54 (outstanding bonds) and E-1 56 (Table B-4)"),
                    "note": (f"Interest only: $315,500,000 at 6.395% (${i1:,.2f}) plus $12,500,000 at 6.145% (${i2:,.2f}), the two 2038 and 2040 term bonds that remain "
                             "(par $328,000,000, matching the outstanding table). No principal falls due before the 2038 sinking fund installments. Check: 35% federal subsidy x "
                             f"(1 - 5.7% sequestration) = ${subsidy_check:,.0f}, and the 2026CD OS prints $6,912 thousand (Table B-4). Coupons from the 2010 OS cover (raw/bonds)."),
                    "why": "This is one old bond's yearly interest to the people who lent the money. It is fixed by the bond's printed rate, and it is already as small as the bond itself."}],
        "residual": {"name": "Other bonds (not printed one by one)",
                     "why": "Several older bonds share this line (parts of 2016D to 2016G, 2017A to 2017C and 2018C), and the public statements we found do not print each bond's share for 2026."},
        "source": SRC_OS,
        "note": "Series 2010B was taken out of the combined 'other' column because its terms are printed and its interest checks against the printed federal subsidy.",
        "side": [{"kind": "pfc_applied", "label": "Passenger facility charges pledged to pay part of 2026 O'Hare airport bond debt service (not split by bond)",
                  "amount": 111_658_000, "period": "2026", "basis": "published", "source": dict(SRC_OS, page="E-1 56 (Table B-4)")}]})

    # ---- 6: Finance General O'Hare 0140 and 9047: side info
    splits.append({
        "target": {"by": "ordinance_line", "fund": "0740", "dept": "99", "authority": "2005", "account": "0140"},
        "expect_amount": ORD_FG_0140, "mode": "side_only",
        "note": ("No public record says who is paid from this line. The payments file has no fund or account, and the Finance General payments cannot be told apart by airport. "
                 "The budget book shows the line was $91.6M in 2025 and $53.0M was actually spent in 2024."),
        "side": [{"kind": "prior_year_budget", "label": "2025 appropriation, same line", "amount": 91_593_627, "period": "2025", "basis": "published",
                  "source": dict(SRC_BOOK, page="585 (PDF 594)")},
                 {"kind": "prior_year_actual", "label": "2024 actual spending, same line", "amount": 53_003_859, "period": "2024", "basis": "actual",
                  "source": dict(SRC_BOOK, page="585 (PDF 594)")}]})
    cap = []
    for lab, amt in (("TAP Phase 1 (ORDNext terminals, tunnel, parking), escalated", 9_826_217_000), ("Pre-approved capital improvement projects, escalated", 1_925_810_000),
                     ("Pre-approved allowances and infrastructure reliability, escalated", 768_157_000), ("Ongoing projects from earlier agreements", 209_218_000),
                     ("New approvals since the 2018 airline agreement", 1_493_760_000), ("Proposed projects", 403_952_000),
                     ("Total airport capital program, escalated dollars", 14_627_114_000)):
        cap.append({"kind": "capital_program", "label": lab, "amount": amt, "period": "program to 2035", "basis": "published",
                    "source": dict(SRC_OS, page="E-19 (Table S-1)")})
    for lab, amt in (("Concourse D (Satellite 1)", 730_000_000), ("Concourse E (Satellite 2), north and south", 527_000_000),
                     ("Terminal 2 redevelopment: O'Hare Global Terminal and Concourse", 1_830_000_000), ("Baggage handling equipment", 825_000_000),
                     ("Consolidated people mover, pedestrian and utility tunnel (phase 1)", 665_000_000), ("Terminal 5 repurposing and core expansion", 250_000_000),
                     ("Central detention basin elimination and south basin expansion", 223_000_000), ("Terminal 5 landside and parking, stage 1 and 2", 222_000_000),
                     ("TAP phase 1 utilities allowance", 493_000_000), ("Western parking and screening allowance", 215_000_000)):
        cap.append({"kind": "capital_project_2018", "label": lab + " (2018 dollars, airline agreement exhibit L)", "amount": amt, "period": "approved 2018", "basis": "published",
                    "source": dict(SRC_OS, page="E-43 (Table 2-1)")})
    splits.append({
        "target": {"by": "ordinance_line", "fund": "0740", "dept": "99", "authority": "2005", "account": "9047"},
        "expect_amount": ORD_FG_9047, "mode": "side_only",
        "note": ("This is the only O'Hare capital line in the ordinance, and it is small. O'Hare's building program (O'Hare 21) is not appropriated line by line: it is paid with "
                 "airport bond proceeds and airline fees under the airline agreement. The side list shows the program's published cost tables."),
        "side": cap + [{"kind": "prior_year_actual", "label": "2024 actual spending, same line", "amount": 1_376_213, "period": "2024", "basis": "actual",
                        "source": dict(SRC_BOOK, page="586 (PDF 595)")}]})

    meta = {"author": "deep: O'Hare & Midway (squid)", "built_by": "scripts/aviation_build.py",
            "description": ("Aviation detail: named FAA grants under the O'Hare and Midway AIP reserve lines, FY2026 awards and the Terminal 3 construction manager "
                            "under the AIP professional services lines, Series 2010B interest out of the O'Hare bond residual, capital program and prior-year side info. "
                            "See research/aviation.md."),
            "awards_pulled": json.load(open(AWARDS_CACHE)).get("pulled") if os.path.exists(AWARDS_CACHE) else None,
            "faa_workbook_mismatches": audits}
    json.dump({"meta": meta, "splits": splits}, open(OUT, "w"), indent=1)
    print(f"wrote {OUT}: {len(splits)} splits")
    for s in splits:
        t = s["target"]
        n = len(s.get("pieces", []))
        tot = sum(p["amount"] for p in s.get("pieces", []))
        print(f"  {t.get('account') or t.get('id')[-40:]:>10} {t.get('authority', ''):>5} expect {s['expect_amount']:>14,.2f}  pieces {n:>2} total {tot:>14,.2f}  mode {s['mode']}")
    if audits:
        print("FAA workbook vs USASpending mismatches:", audits)


if __name__ == "__main__":
    main()
